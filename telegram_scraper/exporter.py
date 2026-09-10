"""
Exporters for saving scraped Telegram posts to JSON, Excel (.xlsx), and CSV.
Includes media download support with polite rate-limiting.
"""

import os
import json
import csv
import logging
import asyncio
from pathlib import Path
from typing import List, Optional
import httpx

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from .models import ScrapeResult, TelegramPost
from .utils import sanitize_filename

logger = logging.getLogger(__name__)


class DataExporter:
    """Handles exporting ScrapeResult data to various file formats and downloading media."""

    def __init__(
        self,
        result: ScrapeResult,
        output_dir: str = "downloads",
        proxy: Optional[str] = None,
    ):
        self.result = result
        self.output_dir = Path(output_dir)
        self.output_dir.mkdir(parents=True, exist_ok=True)
        self.channel_name = result.channel.username
        self.proxy = proxy

    def export_json(self, filename: Optional[str] = None) -> str:
        """Export full structured data to formatted JSON."""
        if not filename:
            filename = f"{self.channel_name}_posts.json"

        file_path = self.output_dir / filename
        data = self.result.model_dump()

        with open(file_path, "w", encoding="utf-8") as f:
            json.dump(data, f, ensure_ascii=False, indent=2)

        logger.info("Exported JSON to: %s", file_path)
        return str(file_path)

    def export_csv(self, filename: Optional[str] = None) -> str:
        """
        Export posts to CSV with UTF-8 BOM encoding for complete Persian/Arabic Excel compatibility.
        """
        if not filename:
            filename = f"{self.channel_name}_posts.csv"

        file_path = self.output_dir / filename

        fieldnames = [
            "post_id",
            "date_iso",
            "date_formatted",
            "views",
            "views_numeric",
            "text",
            "reactions",
            "reactions_count",
            "is_pinned",
            "is_forwarded",
            "forward_from",
            "media_types",
            "media_urls",
            "link_preview_title",
            "link_preview_url",
            "post_url",
        ]

        with open(file_path, "w", encoding="utf-8-sig", newline="") as f:
            writer = csv.DictWriter(f, fieldnames=fieldnames)
            writer.writeheader()

            for post in self.result.posts:
                reactions_str = ", ".join(f"{r.emoji} {r.count}" for r in post.reactions)
                forward_str = post.forward_info.from_name if post.forward_info else ""
                media_types = ", ".join(set(m.type for m in post.media))
                media_urls = " | ".join(m.url or m.thumbnail_url or "" for m in post.media if m.url or m.thumbnail_url)

                writer.writerow(
                    {
                        "post_id": post.id,
                        "date_iso": post.date_iso or "",
                        "date_formatted": post.date_formatted or "",
                        "views": post.views or "",
                        "views_numeric": post.views_numeric or "",
                        "text": post.text,
                        "reactions": reactions_str,
                        "reactions_count": post.total_reactions_count,
                        "is_pinned": "Yes" if post.is_pinned else "No",
                        "is_forwarded": "Yes" if post.is_forwarded else "No",
                        "forward_from": forward_str,
                        "media_types": media_types,
                        "media_urls": media_urls,
                        "link_preview_title": post.link_preview.title if post.link_preview else "",
                        "link_preview_url": post.link_preview.url if post.link_preview else "",
                        "post_url": post.url,
                    }
                )

        logger.info("Exported CSV to: %s", file_path)
        return str(file_path)

    def export_excel(self, filename: Optional[str] = None) -> str:
        """
        Export posts and channel info to a styled Excel (.xlsx) file.
        """
        if not filename:
            filename = f"{self.channel_name}_posts.xlsx"

        file_path = self.output_dir / filename

        wb = Workbook()

        # ----------------- Sheet 1: Posts -----------------
        ws_posts = wb.active
        ws_posts.title = "Posts"

        headers = [
            "شناسه پست (ID)",
            "تاریخ و زمان",
            "متن پیام",
            "تعداد بازدید (Views)",
            "مجموع واکنش‌ها",
            "واکنش‌ها (Reactions)",
            "پین شده؟",
            "فوروارد از",
            "نوع رسانه",
            "لینک رسانه‌ها",
            "پیش‌نمایش لینک",
            "لینک مستقیم تلگرام",
        ]

        ws_posts.append(headers)

        # Style headers
        header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        thin_border = Border(
            left=Side(style="thin", color="D3D3D3"),
            right=Side(style="thin", color="D3D3D3"),
            top=Side(style="thin", color="D3D3D3"),
            bottom=Side(style="thin", color="D3D3D3"),
        )

        for col_num in range(1, len(headers) + 1):
            cell = ws_posts.cell(row=1, column=col_num)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        # Populate rows
        for row_idx, post in enumerate(self.result.posts, start=2):
            reactions_str = " | ".join(f"{r.emoji} {r.count}" for r in post.reactions)
            forward_str = post.forward_info.from_name if post.forward_info else ""
            media_types = ", ".join(set(m.type for m in post.media))
            media_urls = "\n".join(m.url or m.thumbnail_url or "" for m in post.media if m.url or m.thumbnail_url)
            lp_str = (
                f"{post.link_preview.title or ''} ({post.link_preview.url or ''})"
                if post.link_preview
                else ""
            )

            row_values = [
                post.id,
                post.date_iso or post.date_formatted or "",
                post.text,
                post.views or "",
                post.total_reactions_count,
                reactions_str,
                "بله" if post.is_pinned else "خیر",
                forward_str,
                media_types,
                media_urls,
                lp_str,
                post.url,
            ]
            ws_posts.append(row_values)

            # Cell formatting
            for col_num in range(1, len(headers) + 1):
                cell = ws_posts.cell(row=row_idx, column=col_num)
                cell.border = thin_border
                cell.alignment = Alignment(vertical="top", wrap_text=(col_num in (3, 10, 11)))

        # Adjust column widths
        for col in ws_posts.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val = str(cell.value or "")
                # Limit line length estimation
                first_line = val.split("\n")[0] if "\n" in val else val
                if len(first_line) > max_len:
                    max_len = len(first_line)
            ws_posts.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 45)

        # ----------------- Sheet 2: Channel Info -----------------
        ws_info = wb.create_sheet(title="Channel Info")
        ws_info.append(["مشخصه", "مقدار"])

        for col_num in (1, 2):
            c = ws_info.cell(row=1, column=col_num)
            c.fill = header_fill
            c.font = header_font
            c.alignment = Alignment(horizontal="center", vertical="center")

        info_rows = [
            ("شناسه کانال (Username)", f"@{self.result.channel.username}"),
            ("نام کانال (Title)", self.result.channel.title or ""),
            ("تعداد اعضا / مشترکین", self.result.channel.subscribers or ""),
            ("تأیید شده (Verified)", "بله" if self.result.channel.is_verified else "خیر"),
            ("لینک آواتار", self.result.channel.avatar_url or ""),
            ("توضیحات کانال (Bio)", self.result.channel.description or ""),
            ("تاریخ اسکرپ", self.result.scraped_at),
            ("موتور استفاده شده", self.result.engine_used),
            ("تعداد پست‌های استخراج‌شده", str(self.result.total_scraped)),
        ]

        for k, v in info_rows:
            ws_info.append([k, v])

        for row_idx in range(2, len(info_rows) + 2):
            for col_num in (1, 2):
                c = ws_info.cell(row=row_idx, column=col_num)
                c.border = thin_border
                c.alignment = Alignment(vertical="top", wrap_text=True)

        ws_info.column_dimensions["A"].width = 28
        ws_info.column_dimensions["B"].width = 50

        wb.save(file_path)
        logger.info("Exported Excel to: %s", file_path)
        return str(file_path)

    async def download_media_files(self, max_concurrent: int = 4) -> int:
        """
        Download media files (photos, thumbnails) into an organized local folder.
        """
        media_dir = self.output_dir / f"{self.channel_name}_media"
        media_dir.mkdir(parents=True, exist_ok=True)

        download_tasks = []
        for post in self.result.posts:
            for idx, media_item in enumerate(post.media):
                target_url = media_item.url or media_item.thumbnail_url
                if target_url and target_url.startswith("http"):
                    ext = ".jpg"
                    if ".mp4" in target_url:
                        ext = ".mp4"
                    elif ".ogg" in target_url:
                        ext = ".ogg"
                    elif ".webp" in target_url:
                        ext = ".webp"

                    dest_file = media_dir / f"post_{post.id}_{idx + 1}{ext}"
                    download_tasks.append((media_item, target_url, dest_file))

        if not download_tasks:
            return 0

        semaphore = asyncio.Semaphore(max_concurrent)
        downloaded_count = 0

        headers = {
            "User-Agent": (
                "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
                "(KHTML, like Gecko) Chrome/126.0.0.0 Safari/537.36"
            ),
            "Referer": "https://t.me/",
        }

        async with httpx.AsyncClient(
            proxy=self.proxy,
            headers=headers,
            follow_redirects=True,
            timeout=15.0,
        ) as client:
            async def _download(media_item, url, dest_path):
                nonlocal downloaded_count
                async with semaphore:
                    try:
                        r = await client.get(url)
                        if r.status_code == 200:
                            dest_path.write_bytes(r.content)
                            media_item.local_path = str(dest_path)
                            downloaded_count += 1
                    except Exception as e:
                        logger.warning("Failed to download media %s: %s", url, e)

            await asyncio.gather(*[_download(m, u, d) for m, u, d in download_tasks])

        return downloaded_count
